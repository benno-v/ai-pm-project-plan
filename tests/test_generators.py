#!/usr/bin/env python3
"""
Smoke tests for the four project-plan-generator scripts.

Generates every format from the sample plan into a temp directory and
checks: the file exists and is non-empty, and — for the three formats that
can be inspected without extra tooling — that milestones actually render as
markers on the schedule (diamonds in the SVG Gantt / PPTX slides, ``◆``
cells in the XLSX Gantt sheet), not just on their own dedicated slide/sheet.

Run with:
    python -m unittest tests/test_generators.py -v
"""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = REPO_ROOT / "skills" / "project-plan-generator" / "scripts"
SAMPLE_PLAN = REPO_ROOT / "assets" / "sample-outputs" / "sample-plan.json"


def _run(script: str, plan: Path, out: Path, *extra_args: str) -> subprocess.CompletedProcess:
    cmd = [sys.executable, str(SCRIPTS_DIR / script), "--plan", str(plan), "--out", str(out), *extra_args]
    return subprocess.run(cmd, capture_output=True, text=True)


class GeneratorSmokeTests(unittest.TestCase):
    """Generate all four formats from the sample plan into a temp dir."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.tmpdir = tempfile.TemporaryDirectory()
        cls.out_dir = Path(cls.tmpdir.name)
        cls.plan = json.loads(SAMPLE_PLAN.read_text())
        cls.n_milestones = len(cls.plan["milestones"])

        cls.html_path = cls.out_dir / "plan.html"
        cls.pptx_path = cls.out_dir / "plan.pptx"
        cls.xlsx_path = cls.out_dir / "plan.xlsx"
        cls.pdf_path = cls.out_dir / "plan.pdf"

        cls.html_result = _run("generate_html.py", SAMPLE_PLAN, cls.html_path)
        cls.pptx_result = _run("generate_pptx.py", SAMPLE_PLAN, cls.pptx_path)
        cls.xlsx_result = _run("generate_xlsx.py", SAMPLE_PLAN, cls.xlsx_path)
        cls.pdf_result = _run("generate_pdf.py", SAMPLE_PLAN, cls.pdf_path)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.tmpdir.cleanup()

    # -- HTML -----------------------------------------------------------

    def test_html_generates(self) -> None:
        self.assertEqual(self.html_result.returncode, 0, self.html_result.stderr)
        self.assertTrue(self.html_path.exists())
        self.assertGreater(self.html_path.stat().st_size, 0)

    def test_html_has_milestone_diamonds_on_the_gantt(self) -> None:
        html = self.html_path.read_text(encoding="utf-8")
        diamond_count = html.count('class="ms-diamond"')
        self.assertEqual(
            diamond_count, self.n_milestones,
            "expected one milestone diamond on the Gantt per in-range milestone",
        )

    def test_html_has_no_google_fonts_dependency_by_default(self) -> None:
        html = self.html_path.read_text(encoding="utf-8")
        self.assertNotIn("fonts.googleapis.com", html)
        self.assertNotIn("fonts.gstatic.com", html)

    def test_html_has_print_stylesheet(self) -> None:
        html = self.html_path.read_text(encoding="utf-8")
        self.assertIn("@media print", html)
        self.assertIn("@page", html)

    # -- PPTX -------------------------------------------------------------

    def test_pptx_generates(self) -> None:
        self.assertEqual(self.pptx_result.returncode, 0, self.pptx_result.stderr)
        self.assertTrue(self.pptx_path.exists())
        self.assertGreater(self.pptx_path.stat().st_size, 0)

    def test_pptx_has_milestone_diamonds_on_the_schedule_slides(self) -> None:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

        prs = Presentation(str(self.pptx_path))
        diamonds_per_slide = []
        for slide in prs.slides:
            count = 0
            for shape in slide.shapes:
                if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
                    continue
                if shape.auto_shape_type == MSO_SHAPE.DIAMOND:
                    count += 1
            if count:
                diamonds_per_slide.append(count)

        # Timeline overview + Detail Gantt slides each carry a milestones row.
        self.assertEqual(
            len(diamonds_per_slide), 2,
            "expected exactly two Gantt-bearing slides to carry milestone diamonds",
        )
        for count in diamonds_per_slide:
            self.assertEqual(count, self.n_milestones)

    # -- XLSX -------------------------------------------------------------

    def test_xlsx_generates(self) -> None:
        self.assertEqual(self.xlsx_result.returncode, 0, self.xlsx_result.stderr)
        self.assertTrue(self.xlsx_path.exists())
        self.assertGreater(self.xlsx_path.stat().st_size, 0)

    def test_xlsx_has_milestone_diamond_markers_on_the_gantt_sheet(self) -> None:
        from openpyxl import load_workbook

        wb = load_workbook(str(self.xlsx_path))
        self.assertIn("Gantt", wb.sheetnames)
        ws = wb["Gantt"]
        diamond_count = sum(
            1 for row in ws.iter_rows() for cell in row if cell.value == "◆"
        )
        self.assertEqual(diamond_count, self.n_milestones)

    # -- PDF ----------------------------------------------------------------

    def test_pdf_generates(self) -> None:
        if self.pdf_result.returncode != 0:
            self.skipTest(
                "no PDF renderer available in this environment "
                f"(weasyprint/chrome/libreoffice): {self.pdf_result.stderr.strip()}"
            )
        self.assertTrue(self.pdf_path.exists())
        self.assertGreater(self.pdf_path.stat().st_size, 0)
        with open(self.pdf_path, "rb") as f:
            self.assertEqual(f.read(5), b"%PDF-")


if __name__ == "__main__":
    unittest.main()
